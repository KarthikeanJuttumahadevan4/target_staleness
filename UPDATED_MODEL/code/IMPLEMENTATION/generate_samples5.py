#!/usr/bin/env python3
import os
from pathlib import Path
from datetime import datetime, timedelta
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from openpyxl import Workbook
from docx import Document
from pptx import Presentation

base = Path("samples5")
os.makedirs(base, exist_ok=True)

def touch(path, days_ago):
    ts = (datetime.now() - timedelta(days=days_ago)).timestamp()
    os.utime(path, (ts, ts))

# PDFs
def make_pdf(path, lines):
    c = canvas.Canvas(str(path), pagesize=letter)
    y = 750
    for line in lines:
        c.drawString(72, y, line)
        y -= 15
    c.save()

make_pdf(base/"invoices/INV-2024-001.pdf", ["Invoice INV-2024-001","Date: 2024-09-01","Customer: Acme Corp","Product X price is 100","Total: $1000"])
touch(base/"invoices/INV-2024-001.pdf", 30)

make_pdf(base/"invoices/INV-2024-002-adobe-squashed.pdf", ["Invoice INV-2024-002","Date: 2024-09-15","Customer: Acme Corp","Product X price is 110","Total: $1100"])
touch(base/"invoices/INV-2024-002-adobe-squashed.pdf", 20)

make_pdf(base/"invoices/INV-2024-003-ocr-scanned.pdf", ["Scanned Invoice","Date: 2024-10-01","Product X price is 120","Total: $1200"])
touch(base/"invoices/INV-2024-003-ocr-scanned.pdf", 5)

make_pdf(base/"reports/Q3-2024-Financial-Report.pdf", ["Q3 2024 Financial Report","Revenue: $500000","Headcount: 150","Budget Q3: $450000"])
touch(base/"reports/Q3-2024-Financial-Report.pdf", 25)

make_pdf(base/"reports/Q3-2024-Financial-Report-adobe-compressed.pdf", ["Q3 2024 Financial Report Revised","Revenue: $520000","Headcount: 155","Budget Q3: $460000"])
touch(base/"reports/Q3-2024-Financial-Report-adobe-compressed.pdf", 10)

make_pdf(base/"documents/Policy-Update-2024-09-01.pdf", ["Policy Update 2024-09-01","Remote work allowed 3 days per week"])
touch(base/"documents/Policy-Update-2024-09-01.pdf", 40)

make_pdf(base/"scanned/Receipt-Scan-2024-07-22-OCR.pdf", ["Receipt","Date: 2024-07-22","Amount: $250"])
touch(base/"scanned/Receipt-Scan-2024-07-22-OCR.pdf", 90)

make_pdf(base/"scanned/Contract-Scan-2024-08-15-OCR.pdf", ["Contract Scan","Contract value: $50000","Effective: 2024-08-15"])
touch(base/"scanned/Contract-Scan-2024-08-15-OCR.pdf", 80)

# Excel
wb = Workbook()
ws = wb.active
ws.title = "Budget"
ws.append(["Category","Q3 Budget","Q3 Actual"])
ws.append(["Sales",450000,470000])
ws.append(["R&D",120000,115000])
wb.save(base/"spreadsheets/Budget-Q3-2024.xlsx")
touch(base/"spreadsheets/Budget-Q3-2024.xlsx", 22)

wb2 = Workbook()
ws2 = wb2.active
ws2.title = "Headcount"
ws2.append(["Department","Headcount","Hire Date"])
ws2.append(["Engineering",80,"2024-09-01"])
ws2.append(["Sales",45,"2024-09-01"])
wb2.save(base/"spreadsheets/Headcount-2024-Q3.xlsx")
touch(base/"spreadsheets/Headcount-2024-Q3.xlsx", 18)

# Word
doc = Document()
doc.add_heading("Employee Handbook 2024",0)
doc.add_paragraph("Company policy: Remote work allowed 3 days per week.")
doc.add_paragraph("Headcount target: 150 employees.")
doc.save(base/"documents/Employee-Handbook-2024.docx")
touch(base/"documents/Employee-Handbook-2024.docx", 60)

doc2 = Document()
doc2.add_heading("Meeting Minutes 2024-09-10",0)
doc2.add_paragraph("Attendees: Alice, Bob")
doc2.add_paragraph("Decision: Product X price set to 120")
doc2.save(base/"documents/Meeting-Minutes-2024-09-10.docx")
touch(base/"documents/Meeting-Minutes-2024-09-10.docx", 7)

doc3 = Document()
doc3.add_heading("NDA AcmeCorp 2024",0)
doc3.add_paragraph("Contract value: $50000")
doc3.save(base/"contracts/NDA-AcmeCorp-2024.docx")
touch(base/"contracts/NDA-AcmeCorp-2024.docx", 35)

doc4 = Document()
doc4.add_heading("SOW Project Phoenix 2024",0)
doc4.add_paragraph("Budget: $250000")
doc4.add_paragraph("Headcount: 20")
doc4.save(base/"contracts/SOW-ProjectPhoenix-2024.docx")
touch(base/"contracts/SOW-ProjectPhoenix-2024.docx", 15)

# PowerPoint
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Sales Deck 2024-09"
slide.placeholders[1].text = "Product X price is 120\nRevenue $520000"
prs.save(base/"presentations/Sales-Deck-2024-09.pptx")
touch(base/"presentations/Sales-Deck-2024-09.pptx", 12)

prs2 = Presentation()
slide2 = prs2.slides.add_slide(prs2.slide_layouts[0])
slide2.shapes.title.text = "Q3 Board Presentation"
slide2.placeholders[1].text = "Headcount 155\nBudget Q3 $460000"
prs2.save(base/"presentations/Q3-Board-Presentation.pptx")
touch(base/"presentations/Q3-Board-Presentation.pptx", 9)

# Text
txt = base/"demo/price.txt"
txt.write_text("Product X price is 130")
touch(txt, 1)

print("Generation complete")
